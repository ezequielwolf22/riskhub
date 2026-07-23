"""Extraccion de texto, chunking y gestion de documentos IA.

Seguridad en reposo:
  Los archivos subidos se cifran con Fernet (AES-128-CBC + HMAC-SHA256) antes
  de escribirse en disco. La clave se deriva del SECRET_KEY del servidor usando
  SHA-256 -> base64url. Los archivos anteriores a esta version (no cifrados) se
  leen correctamente por compatibilidad hacia atras (InvalidToken -> raw bytes).
"""
import base64
import hashlib
import io
import re
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session
from sqlalchemy import text

from app.config import settings
from app.models import AiDocument, AiDocumentChunk, AiDocumentStatus

# Directorio de almacenamiento de documentos
# FIX: comprobar el punto de montaje persistente (/srv/data) en vez de la
# subcarpeta "documents" — esta ultima nunca existe en el primer arranque de
# un volumen nuevo, asi que el chequeo anterior caia SIEMPRE al fallback local
# (efimero, se pierde en cada `docker build` + recreate del contenedor).
_DOC_ROOT = Path("/srv/data/documents")
if not _DOC_ROOT.parent.exists():
    _DOC_ROOT = Path(__file__).parent.parent.parent / "data" / "documents"
_DOC_ROOT.mkdir(parents=True, exist_ok=True)

CHUNK_SIZE = 600
CHUNK_OVERLAP = 100


# ---------- Cifrado en reposo ----------

def _fernet_key() -> bytes:
    """Deriva clave Fernet de 32 bytes a partir del SECRET_KEY de la app."""
    return base64.urlsafe_b64encode(
        hashlib.sha256(settings.secret_key.encode()).digest()
    )


def encrypt_doc(data: bytes) -> bytes:
    """Cifra los bytes de un documento con Fernet antes de escribir en disco."""
    from cryptography.fernet import Fernet
    return Fernet(_fernet_key()).encrypt(data)


def decrypt_doc(data: bytes) -> bytes:
    """Descifra los bytes de un documento leido del disco.

    Si el archivo no esta cifrado (version anterior), devuelve los bytes sin
    modificar para mantener compatibilidad con documentos pre-cifrado.
    """
    from cryptography.fernet import Fernet, InvalidToken
    try:
        return Fernet(_fernet_key()).decrypt(data)
    except (InvalidToken, ValueError, Exception):
        # Archivo pre-cifrado: devolver raw para compatibilidad hacia atras
        return data


def save_document_file(data: bytes, filename: str) -> None:
    """Escribe los bytes del documento en disco cifrados con Fernet."""
    dest = _DOC_ROOT / filename
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(encrypt_doc(data))


# ---------- helpers ----------

def doc_path(filename: str) -> Path:
    return _DOC_ROOT / filename


def _extract_text_pdf(data: bytes) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                parts.append(t)
        return "\n".join(parts)
    except Exception as e:
        raise ValueError(f"Error extrayendo PDF: {e}")


def _extract_text_docx(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        raise ValueError(f"Error extrayendo DOCX: {e}")


def _extract_text_plain(data: bytes) -> str:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_text_xlsx(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Hoja: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip()]
                if cells:
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)
    except Exception as e:
        raise ValueError(f"Error extrayendo XLSX: {e}")


def extract_text(data: bytes, mime_type: str) -> str:
    if "pdf" in mime_type:
        return _extract_text_pdf(data)
    if "wordprocessingml" in mime_type or "docx" in mime_type:
        return _extract_text_docx(data)
    if "spreadsheetml" in mime_type or "xlsx" in mime_type:
        return _extract_text_xlsx(data)
    return _extract_text_plain(data)


def chunk_text(full_text: str) -> list[str]:
    """Divide el texto en fragmentos de CHUNK_SIZE con solapamiento CHUNK_OVERLAP."""
    full_text = re.sub(r"\s+", " ", full_text).strip()
    chunks = []
    start = 0
    while start < len(full_text):
        end = start + CHUNK_SIZE
        chunks.append(full_text[start:end])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return [c for c in chunks if len(c.strip()) > 20]


def _delete_fts_for_doc(db: Session, doc_id: int) -> None:
    """Elimina del indice FTS5 todos los chunks de un documento."""
    chunk_ids = [
        r[0] for r in db.query(AiDocumentChunk.id).filter_by(document_id=doc_id).all()
    ]
    for cid in chunk_ids:
        try:
            db.execute(text("DELETE FROM ai_chunks_fts WHERE rowid = :cid"), {"cid": cid})
        except Exception:
            pass  # FTS puede no existir aun


def _index_text(db: Session, doc: AiDocument, raw_text: str) -> int:
    """Chunking + indexado FTS5 del texto de un documento (reemplaza chunks previos)."""
    chunks = chunk_text(raw_text)

    _delete_fts_for_doc(db, doc.id)
    db.query(AiDocumentChunk).filter_by(document_id=doc.id).delete()

    for idx, chunk_content in enumerate(chunks):
        chunk = AiDocumentChunk(
            document_id=doc.id,
            chunk_index=idx,
            content=chunk_content,
        )
        db.add(chunk)
        db.flush()  # necesario para obtener chunk.id
        try:
            db.execute(
                text("INSERT INTO ai_chunks_fts(rowid, content) VALUES (:rowid, :content)"),
                {"rowid": chunk.id, "content": chunk_content},
            )
        except Exception:
            pass  # FTS puede no estar disponible

    doc.chunk_count = len(chunks)
    doc.status = AiDocumentStatus.INDEXED
    doc.processed_at = datetime.now(timezone.utc)
    return len(chunks)


def process_document(db: Session, doc_id: int) -> None:
    """Extrae texto, genera chunks y los indexa en FTS5."""
    doc = db.query(AiDocument).filter_by(id=doc_id).first()
    if not doc:
        return

    doc.status = AiDocumentStatus.PROCESSING
    db.commit()

    try:
        # decrypt_doc soporta tanto archivos cifrados (nuevos) como no cifrados (legacy)
        file_data = decrypt_doc(doc_path(doc.filename).read_bytes())
        raw_text = extract_text(file_data, doc.mime_type or "text/plain")
        _index_text(db, doc, raw_text)
        db.commit()

        # Generar embeddings semanticos si hay Voyage API key configurada
        _try_embed_document(db, doc_id, doc.organization_id)

    except Exception as e:
        doc.status = AiDocumentStatus.ERROR
        doc.error_message = str(e)[:500]
        db.commit()
        raise


def describe_document_with_vision(doc_id: int) -> bool:
    """Transcribe con Claude Vision documentos sin texto extraible.

    Cubre imagenes (capturas, diagramas, actas fotografiadas) y PDFs
    escaneados sin capa de texto: antes quedaban INDEXED con 0 chunks,
    invisibles para el RAG y el analisis ISMS. Crea su propia sesion.
    Devuelve True si genero chunks.
    """
    import base64 as _b64
    import logging
    from app.database import SessionLocal
    logger = logging.getLogger("riskhub.doc")

    db = SessionLocal()
    try:
        doc = db.get(AiDocument, doc_id)
        if not doc or not doc.filename:
            return False
        mime = doc.mime_type or ""

        from app.services.model_registry import get_api_key, get_model
        api_key = get_api_key(db, doc.organization_id)
        if not api_key:
            return False

        raw = decrypt_doc(doc_path(doc.filename).read_bytes())
        if mime.startswith("image/"):
            if mime == "image/svg+xml" or len(raw) > 5 * 1024 * 1024:
                return False
            media = mime if mime in ("image/png", "image/jpeg", "image/gif", "image/webp") else "image/png"
            block = {"type": "image",
                     "source": {"type": "base64", "media_type": media,
                                "data": _b64.standard_b64encode(raw).decode()}}
        elif "pdf" in mime:
            if len(raw) > 20 * 1024 * 1024:
                return False
            block = {"type": "document",
                     "source": {"type": "base64", "media_type": "application/pdf",
                                "data": _b64.standard_b64encode(raw).decode()}}
        else:
            return False

        from app.services.claude_client import create_message
        msg = create_message(
            api_key,
            model=get_model(db, doc.organization_id, tier="fast"),
            max_tokens=8192,
            org_id=doc.organization_id, call_type="document_vision",
            system=(
                "Transcribe fielmente el contenido del documento o imagen al espanol. "
                "Incluye todo el texto legible, tablas como lineas 'campo: valor', y "
                "describe brevemente los elementos visuales relevantes (sellos, firmas, "
                "graficos, capturas de configuracion). No inventes contenido ilegible: "
                "marca [ilegible]. Devuelve solo la transcripcion, sin comentarios."
            ),
            messages=[{"role": "user", "content": [
                block,
                {"type": "text", "text": f"Documento: {doc.original_name}"},
            ]}],
        )
        transcription = msg.content[0].text if msg.content else ""
        if len(transcription.strip()) < 30:
            return False

        n = _index_text(db, doc, f"[Transcripcion automatica via Vision]\n{transcription}")
        doc.error_message = None
        db.commit()
        _try_embed_document(db, doc_id, doc.organization_id)
        logger.info("Vision: doc %d transcrito (%d chunks)", doc_id, n)
        return n > 0
    except Exception as e:
        import logging
        logging.getLogger("riskhub.doc").warning("Vision fallo para doc %d: %s", doc_id, e)
        return False
    finally:
        db.close()


def _try_embed_document(db: Session, doc_id: int, organization_id: int | None) -> None:
    """Genera embeddings Voyage AI para el documento si la key esta configurada.

    Falla silenciosamente — los embeddings son una mejora opcional sobre FTS5.
    """
    if not organization_id:
        return
    try:
        from app.models import AiConfig
        from app.security import decrypt_secret
        from app.services.embedding_service import embed_document_chunks
        cfg = db.query(AiConfig).filter_by(organization_id=organization_id).first()
        if not cfg or not cfg.voyage_api_key_encrypted:
            return
        voyage_key = decrypt_secret(cfg.voyage_api_key_encrypted)
        embed_document_chunks(db, doc_id, voyage_key)
    except Exception as e:
        import logging
        logging.getLogger("riskhub.doc").warning("Embedding falló para doc %d: %s", doc_id, e)


# Entidades que referencian a un documento. El borrado nunca las destruye:
# son registros del SGSI con vida propia (una politica tiene codigo, version y
# aprobacion; un plan BCP tiene su ciclo). Se desvinculan y se informa de ello.
# (tabla, campo, clave i18n del tipo para el resumen)
_DOC_REFERENCES: list[tuple[str, str, str]] = [
    ("Policy", "source_document_id", "policies"),
    ("BCPPlan", "document_id", "bcp_plans"),
    ("Supplier", "dpa_document_id", "suppliers"),
    ("Supplier", "nda_document_id", "suppliers"),
    ("Supplier", "contract_document_id", "suppliers"),
    ("StrategicInitiative", "source_document_id", "initiatives"),
    ("IngestSourceMap", "document_id", "ingest_maps"),
]


def _reference_queries(db: Session, doc_id: int):
    """Genera (kind, query) por cada referencia viva al documento."""
    import app.models as models
    for model_name, field, kind in _DOC_REFERENCES:
        model = getattr(models, model_name, None)
        if model is None:
            continue
        column = getattr(model, field, None)
        if column is None:
            continue
        yield kind, field, db.query(model).filter(column == doc_id)


def document_references(db: Session, doc_id: int) -> dict[str, int]:
    """Cuenta que registros quedarian desvinculados si se borra el documento.

    Se usa para el preview del borrado y para el resumen que devuelve la API.
    """
    counts: dict[str, int] = {}
    for kind, _field, query in _reference_queries(db, doc_id):
        n = query.count()
        if n:
            counts[kind] = counts.get(kind, 0) + n
    return counts


def delete_document(db: Session, doc: AiDocument) -> dict[str, int]:
    """Elimina documento, sus chunks de BD y FTS5, y el archivo del disco.

    Antes de borrar desvincula las referencias de otras entidades. Sin esto el
    DELETE fallaba con IntegrityError (PRAGMA foreign_keys=ON) en cuanto el
    analisis ISMS habia derivado una Policy del documento — es decir, casi
    siempre. Devuelve el recuento de lo desvinculado.
    """
    doc_id = doc.id
    detached = document_references(db, doc_id)

    for _kind, field, query in _reference_queries(db, doc_id):
        query.update({field: None}, synchronize_session=False)

    # Evidencia derivada del documento (F3): la auto-generada por la ingesta se
    # elimina con el documento (no tiene vida propia); la que un humano vinculo a
    # mano se conserva desvinculada. Se recalcula el riesgo de los controles cuya
    # evidencia desaparece, para que el residual vuelva a subir sin ella.
    ev_detached, ev_impl_ids = _detach_document_evidence(db, doc_id)
    if ev_detached:
        detached["evidence"] = ev_detached

    _delete_fts_for_doc(db, doc_id)
    db.query(AiDocumentChunk).filter_by(document_id=doc_id).delete()
    p = doc_path(doc.filename)
    if p.exists():
        p.unlink(missing_ok=True)
    db.delete(doc)
    db.commit()

    # Recalcular el riesgo de los controles que perdieron evidencia auto-generada.
    if ev_impl_ids:
        try:
            from app.services.risk_recalc_service import recalc_risks_for_impls
            recalc_risks_for_impls(db, list(ev_impl_ids))
        except Exception:
            pass
    return detached


def _detach_document_evidence(db: Session, doc_id: int) -> tuple[int, set[int]]:
    """Elimina la evidencia auto-generada del documento y desvincula la manual.

    Devuelve (numero de evidencias afectadas, ids de ControlImplementation cuyo
    riesgo conviene recalcular porque perdieron evidencia).
    """
    from app.models import Evidence
    src_col = getattr(Evidence, "source_document_id", None)
    if src_col is None:      # esquema antiguo sin la columna: nada que hacer
        return 0, set()
    rows = db.query(Evidence).filter(src_col == doc_id).all()
    if not rows:
        return 0, set()
    impl_ids: set[int] = set()
    affected = 0
    for ev in rows:
        if ev.control_implementation_id:
            impl_ids.add(ev.control_implementation_id)
        if getattr(ev, "auto_generated", False):
            db.delete(ev)
        else:
            ev.source_document_id = None
        affected += 1
    return affected, impl_ids
